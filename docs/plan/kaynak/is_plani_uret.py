# Is plani takvimi uretici: python docs/plan/kaynak/is_plani_uret.py docs/plan/kaynak/is_plani.json docs/plan/EBM_SmartClinical_Is_Plani_Takvimi.docx
import json
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

d = json.load(open(sys.argv[1]))
CIKTI = sys.argv[2]
LACIVERT, TURKUAZ, ACIK, KT_RENK, GRI = "0F2240", "1F8A8A", "BFE3DF", "D9962B", "F2F4F3"
YAZI = RGBColor(0x0F, 0x22, 0x40)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for k in ("left_margin", "right_margin"): setattr(sec, k, Cm(1.5))
sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

def golge(h, renk):
    tcPr = h._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), renk); tcPr.append(s)

def kenar(tablo, renk="C9D1CE"):
    tblPr = tablo._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for y in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{y}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), renk); b.append(e)
    tblPr.append(b)

def yaz(h, metin, boy=9, kalin=False, renk=None, hiza=None):
    h.text = ""; p = h.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
    if hiza: p.alignment = hiza
    r = p.add_run(metin); r.font.size = Pt(boy); r.bold = kalin
    if renk: r.font.color.rgb = renk
    h.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p

def yukseklik(satir, cm):
    trPr = satir._tr.get_or_add_trPr(); e = OxmlElement("w:trHeight")
    e.set(qn("w:val"), str(int(cm * 567))); e.set(qn("w:hRule"), "atLeast"); trPr.append(e)

def genislik(tablo, olculer):
    tablo.autofit = False
    tblPr = tablo._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(olculer) * 567))); tw.set(qn("w:type"), "dxa")
    for e in tblPr.findall(qn("w:tblW")): tblPr.remove(e)
    tblPr.append(tw)
    grid = tablo._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(olculer): gc.set(qn("w:w"), str(int(olculer[i] * 567)))
    for satir in tablo.rows:
        for h, w in zip(satir.cells, olculer): h.width = Cm(w)

def baslik(metin, seviye=1):
    p = doc.add_paragraph(); r = p.add_run(metin); r.bold = True; r.font.color.rgb = YAZI
    r.font.size = Pt(16 if seviye == 1 else 12)
    p.paragraph_format.space_before = Pt(10 if seviye == 1 else 8); p.paragraph_format.space_after = Pt(4)
    return p

# Kapak satırları
p = doc.add_paragraph(); r = p.add_run("EBM SmartClinical"); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = YAZI
p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph(); r = p.add_run("İş Planı ve Zaman Çizelgesi"); r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x1F, 0x8A, 0x8A)
p = doc.add_paragraph()
r = p.add_run("Proje süresi 18 aydır. Plan, ön başvuruda tanımlanan beş ana iş paketine ve ay aralıklarına göre hazırlanmıştır. "
              "Her iş paketi alt faaliyetlere ayrılmış; çıktılar, başarı ölçütleri, kilometre taşları ve riskler aşağıda verilmiştir.")
r.font.size = Pt(10)

# Gantt
KISA = {
 "İP1.1": "Veri envanteri, etik kurul ve KVKK süreçleri",
 "İP1.2": "FHIR/OMOP uyumlu ortak veri modeli",
 "İP1.3": "Kimliksizleştirme aracının gerçek veriyle doğrulanması",
 "İP1.4": "Altın standart referans veri setinin etiketlenmesi",
 "İP2.1": "Temel modelin Türkçe tıbbi derlemle alana uyarlanması",
 "İP2.2": "Klinik araştırma görevleri için talimat ince ayarı",
 "İP2.3": "Kriter yapılandırma ve RAG tabanlı eşleştirme",
 "İP2.4": "Halüsinasyon kontrolü ve nicemlenmiş hastane içi model",
 "İP3.1": "Hiyerarşik Bayesçi hasta kazanım modeli",
 "İP3.2": "Sağkalım tabanlı aktivasyon gecikme riski modeli",
 "İP3.3": "KRI/QTL tabanlı erken uyarı motoru",
 "İP3.4": "Transfer öğrenme, açıklanabilirlik, geriye dönük doğrulama",
 "İP4.1": "Mikroservis platform ve karar destek panelleri",
 "İP4.2": "HBYS ve EDC için FHIR entegrasyon adaptörleri",
 "İP4.3": "Annex 11 / Part 11, IEC 62304, ISO 27001, sızma testi",
 "İP4.4": "Validasyon dosyası (IQ/OQ/PQ, izlenebilirlik)",
 "İP5.1": "Üç pilot hastanede kurulum ve gölge mod",
 "İP5.2": "Manuel ön tarama ile karşılaştırma ve süre ölçümü",
 "İP5.3": "Kabul testleri, regülasyon sınıflandırma, dokümantasyon",
 "İP5.4": "İş modeli, fikri haklar ve yaygınlaştırma planı",
}
baslik("1. Zaman çizelgesi")
AY = 18
t = doc.add_table(rows=1, cols=2 + AY); kenar(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
olc = [1.4, 10.0] + [0.85] * AY
ust = t.rows[0].cells
yaz(ust[0], "Kod", 8, True, RGBColor(0xFF, 0xFF, 0xFF)); golge(ust[0], LACIVERT)
yaz(ust[1], "İş paketi / faaliyet", 8, True, RGBColor(0xFF, 0xFF, 0xFF)); golge(ust[1], LACIVERT)
for a in range(AY):
    yaz(ust[2 + a], str(a + 1), 8, True, RGBColor(0xFF, 0xFF, 0xFF), WD_ALIGN_PARAGRAPH.CENTER); golge(ust[2 + a], LACIVERT)
for ip in d["paketler"]:
    satir_ = t.add_row(); yukseklik(satir_, 0.56); h = satir_.cells
    yaz(h[0], ip["kod"], 8, True); yaz(h[1], ip["baslik"], 8, True)
    golge(h[0], GRI); golge(h[1], GRI)
    for a in range(AY):
        golge(h[2 + a], LACIVERT if ip["bas"] <= a + 1 <= ip["bit"] else "FFFFFF")
    for f in ip["faaliyetler"]:
        satir_ = t.add_row(); yukseklik(satir_, 0.46); h = satir_.cells
        yaz(h[0], f["kod"], 7.5); yaz(h[1], KISA.get(f["kod"], f["ad"]), 7.5)
        for a in range(AY):
            if f["bas"] <= a + 1 <= f["bit"]: golge(h[2 + a], TURKUAZ)
h = t.add_row().cells
yaz(h[0], "KT", 8, True); yaz(h[1], "Kilometre taşları", 8, True); golge(h[0], GRI); golge(h[1], GRI)
aylar = {}
for k in d["kilometre_taslari"]: aylar.setdefault(k["ay"], []).append(k["kod"])
for a in range(AY):
    if a + 1 in aylar:
        yaz(h[2 + a], "/".join([aylar[a + 1][0]] + [x[2:] for x in aylar[a + 1][1:]]), 6.5, True, RGBColor(0xFF, 0xFF, 0xFF), WD_ALIGN_PARAGRAPH.CENTER); golge(h[2 + a], KT_RENK)
genislik(t, olc)
for satir in t.rows:
    for hc in satir.cells:
        for pp in hc.paragraphs:
            pf = pp.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
            if not pp.text:
                rr = pp.add_run(""); rr.font.size = Pt(6)
                pPr = pp._p.get_or_add_pPr(); rpr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "12"); rpr.append(sz); pPr.append(rpr)
mar = OxmlElement("w:tblCellMar")
for y, v in (("top", 20), ("bottom", 20), ("left", 60), ("right", 40)):
    e = OxmlElement(f"w:{y}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
t._tbl.tblPr.append(mar)
# başlık satırı her sayfada tekrar etsin
trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
for metin, renk in (("■ ", "0F2240"), ("İş paketi süresi   ", None), ("■ ", "1F8A8A"), ("Faaliyet süresi   ", None), ("■ ", "D9962B"), ("Kilometre taşı", None)):
    r = p.add_run(metin); r.font.size = Pt(8)
    if renk: r.font.color.rgb = RGBColor.from_string(renk)

# İş paketi ayrıntıları
baslik("2. İş paketleri").paragraph_format.page_break_before = True
for ip in d["paketler"]:
    baslik(f'{ip["kod"]} – {ip["baslik"]} (Ay {ip["bas"]}–{ip["bit"]})', 2).paragraph_format.keep_with_next = True
    t = doc.add_table(rows=0, cols=2); kenar(t)
    def satir(etiket, icerik):
        h = t.add_row().cells; yaz(h[0], etiket, 9, True); golge(h[0], GRI)
        h[1].text = ""
        if isinstance(icerik, list):
            for i, m in enumerate(icerik):
                p = h[1].paragraphs[0] if i == 0 else h[1].add_paragraph()
                p.paragraph_format.space_after = Pt(1); r = p.add_run(m); r.font.size = Pt(9)
        else:
            p = h[1].paragraphs[0]; r = p.add_run(icerik); r.font.size = Pt(9)
    satir("Amaç", ip["amac"])
    satir("Sorumlu", ip["sorumlu"])
    satir("Faaliyetler", [f'{f["kod"]}  {f["ad"]} (Ay {f["bas"]}–{f["bit"]})' for f in ip["faaliyetler"]])
    satir("Çıktılar", ["– " + c for c in ip["ciktilar"]])
    satir("Başarı ölçütü", ip["basari"])
    genislik(t, [3.4, 23.3])
    for i, satir in enumerate(t.rows):
        trPr = satir._tr.get_or_add_trPr(); e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); trPr.append(e)
        if i < len(t.rows) - 1:
            for hc in satir.cells:
                for pp in hc.paragraphs: pp.paragraph_format.keep_with_next = True

# Kilometre taşları
baslik("3. Kilometre taşları").paragraph_format.page_break_before = True
t = doc.add_table(rows=1, cols=4); kenar(t)
for h, m in zip(t.rows[0].cells, ("Kod", "Ay", "Kilometre taşı", "Doğrulama yöntemi")):
    yaz(h, m, 9, True, RGBColor(0xFF, 0xFF, 0xFF)); golge(h, LACIVERT)
for k in d["kilometre_taslari"]:
    h = t.add_row().cells
    yaz(h[0], k["kod"], 9, True); yaz(h[1], str(k["ay"]), 9, hiza=WD_ALIGN_PARAGRAPH.CENTER); yaz(h[2], k["ad"], 9); yaz(h[3], k["dogrulama"], 9)
genislik(t, [1.6, 1.4, 10.2, 13.5])

baslik("4. Riskler ve B planı")
t = doc.add_table(rows=1, cols=4); kenar(t)
for h, m in zip(t.rows[0].cells, ("Risk", "İş paketi", "Olasılık", "Önlem / B planı")):
    yaz(h, m, 9, True, RGBColor(0xFF, 0xFF, 0xFF)); golge(h, LACIVERT)
for k in d["riskler"]:
    h = t.add_row().cells
    yaz(h[0], k["risk"], 9); yaz(h[1], k["ip"], 9, hiza=WD_ALIGN_PARAGRAPH.CENTER); yaz(h[2], k["olasilik"], 9, hiza=WD_ALIGN_PARAGRAPH.CENTER); yaz(h[3], k["onlem"], 9)
genislik(t, [6.2, 1.9, 1.9, 16.7])

doc.save(CIKTI)
print("kaydedildi", CIKTI)
